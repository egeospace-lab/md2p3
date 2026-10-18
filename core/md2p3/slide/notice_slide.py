"""
Notice Slide Class
==================

Represents the notice slide in a presentation. This slide type is
intended for simple notices or separators and does not require
additional inputs beyond layout and slide indices.

Functional Overview
-------------------
- Inherits from the abstract Slide base class.
- Requires only layoutIndex and optional slideIndex.
- Integrates into a PowerPoint presentation by:
  * Selecting the specified slide layout.
  * Adding a new slide based on that layout.
  * Removing all placeholders to ensure a clean notice slide.
- Does not require additional content or placeholders.
- Ensures graceful failure if integration encounters errors.

Contents
--------
class NoticeSlide(Slide)
    Represents the notice slide in a presentation.
"""

from typing import Optional
from pptx import Presentation

from core.md2p3.slide import Slide


class NoticeSlide(Slide):
    """
    NoticeSlide extends Slide but does not add new fields.
    It simply represents a slide that is placed using its layout index.

    Notes
    -----
    - Designed for simple notice or separator slides.
    - No additional content or formatting is applied.
    """

    def __init__(self, layoutIndex: int, slideIndex: Optional[int] = None) -> None:
        """
        Initialise a NoticeSlide object.

        Parameters
        ----------
        layoutIndex : int
            Index of the layout in the target template file.
        slideIndex : Optional[int], default=None
            Position of the slide in the output presentation.

        Raises
        ------
        TypeError
            If layoutIndex or slideIndex are not integers when provided.
        """
        super().__init__(layoutIndex, slideIndex)

    def integrateWith(self, presentation: Presentation) -> None:
        """
        Integrate the notice slide into the presentation.

        Parameters
        ----------
        presentation : pptx.Presentation
            The PowerPoint presentation object to integrate into.

        Notes
        -----
        - Adds a new slide using the specified layout index.
        - Removes all placeholders since the notice slide does not require any.
        - This ensures no extra boxes remain on the slide.
        - Does not insert additional content; the layout itself conveys the notice.

        Examples
        --------
        Add a notice slide to a presentation:
            >>> notice = NoticeSlide(layoutIndex=2)
            >>> notice.integrateWith(presentation)

        Raises
        ------
        RuntimeError
            If integration fails due to an error.
        """
        try:
            # --- Add new slide using the specified layout ---
            slide = self.addSlide(presentation)

            # --- Clean up placeholders ---
            # Remove all placeholders to avoid extra boxes.
            for shape in list(slide.placeholders):
                sp = shape._element
                sp.getparent().remove(sp)

        except Exception as error:
            # Graceful failure: log and raise a controlled exception
            print(f"[Error] Failed to integrate NoticeSlide: {error}")
            raise RuntimeError("Integration of NoticeSlide failed.") from error
