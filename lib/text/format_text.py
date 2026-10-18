"""
Text Formatting Utilities
=========================

This module provides utilities for applying Markdown-style formatting
to text runs in PowerPoint presentations. Supported tokens include:

- Bold: **text**
- Italic: *text* or _text_
- Inline code: `text` (rendered in red Courier font)

Multiple styles can be applied together (e.g., bold + italic, italic + code).
Inline code spans always render literally, but can inherit surrounding styles
(e.g., **`code`** → bold + red Courier).

Functional Overview
-------------------
- Accepts a PowerPoint paragraph object (`pptx.text.text._Paragraph`).
- Parses input text containing Markdown-style tokens using `markdown-it-py`.
- Applies formatting to runs within the paragraph:
  * Bold text enclosed in **double asterisks**.
  * Italic text enclosed in *single asterisks* or _underscores_.
  * Inline code enclosed in backticks (`), rendered in red Courier font.
  * Code spans ignore emphasis markers inside, but can inherit surrounding styles.
- Strips heading markers (#, ##, ###, etc.) automatically.
- Supports nested emphasis (e.g., **bold *italic***).
- Supports nested bullet levels via the `bulletLevel` parameter.
- Falls back to plain text if parsing fails, ensuring resilience.

Contents
--------
formatText(paragraph: _Paragraph, inputText: str, bulletLevel: int = 0) -> None
    Apply Markdown-style tokenisation to a PowerPoint paragraph.
"""

from pptx.text.text import _Paragraph  # type hint for python-pptx paragraph
from pptx.dml.color import RGBColor
from markdown_it import MarkdownIt


def formatText(paragraph: _Paragraph, inputText: str, bulletLevel: int = 0) -> None:
    """
    Apply Markdown-style tokenisation to a PowerPoint paragraph.

    Parameters
    ----------
    paragraph : pptx.text.text._Paragraph
        The PowerPoint paragraph object to which formatted runs will be added.
    inputText : str
        The input text containing Markdown-style tokens.
    bulletLevel : int, default=0
        Bullet depth level (0 = top-level, 1 = nested, etc.).

    Notes
    -----
    - Bold text is enclosed in **double asterisks**.
    - Italic text is enclosed in *single asterisks* or _underscores_.
    - Inline code is enclosed in backticks (`), rendered in red Courier font.
    - Inline code spans ignore emphasis markers inside, but can inherit
      surrounding bold/italic styles.
    - Heading markers (#, ##, ###, etc.) are stripped automatically.
    - Nested emphasis is supported (e.g., **bold *italic***).
    - Multiple styles can be combined (e.g., ***text*** → bold + italic).
    - If parsing fails, the original text is added as a plain run.

    Examples
    --------
    Apply bold and italic formatting:
        >>> formatText(paragraph, "This is **bold** and *italic* text")

    Apply inline code formatting:
        >>> formatText(paragraph, "Use `print()` in Python")

    Apply bold around code:
        >>> formatText(paragraph, "This is **`code`**")

    Apply nested emphasis:
        >>> formatText(paragraph, "This is **bold *italic*** text")

    Apply nested bullet level:
        >>> formatText(paragraph, "- Nested item", bulletLevel=1)
    """
    try:
        if not isinstance(inputText, str):
            raise TypeError("Argument 'inputText' must be a string.")
        if not isinstance(bulletLevel, int):
            raise TypeError("Argument 'bulletLevel' must be an integer.")

        # Set bullet depth level
        paragraph.level = bulletLevel

        # Remove heading markers (e.g., #, ##, ###)
        cleanedText: str = inputText.lstrip("#").strip()

        # Parse tokens with markdown-it
        md = MarkdownIt()
        tokens = md.parse(cleanedText)

        # Walk through parsed tokens
        for token in tokens:
            if token.type == "inline":
                # Inline tokens contain children (text, strong, em, code_inline, etc.)
                styleStack = []
                for child in token.children or []:
                    if child.type in ("strong_open", "em_open"):
                        styleStack.append(child.type)
                        continue
                    elif child.type in ("strong_close", "em_close"):
                        if styleStack:
                            styleStack.pop()
                        continue
                    elif child.type == "code_inline":
                        # Inline code: rendered literally, but can inherit surrounding styles
                        run = paragraph.add_run()
                        run.text = child.content
                        run.font.color.rgb = RGBColor(255, 0, 0)  # red
                        run.font.name = "Courier New"
                        if "strong_open" in styleStack:
                            run.font.bold = True
                        if "em_open" in styleStack:
                            run.font.italic = True
                        continue
                    elif child.type == "text":
                        # Defensive cleanup: strip any leftover markers
                        content = child.content.strip("*_")
                        run = paragraph.add_run()
                        run.text = content
                        if "strong_open" in styleStack:
                            run.font.bold = True
                        if "em_open" in styleStack:
                            run.font.italic = True

            elif token.type == "text":
                # Plain text outside inline container
                run = paragraph.add_run()
                run.text = token.content

    except Exception as error:
        # Safe failure: fall back to plain text
        run = paragraph.add_run()
        run.text = str(inputText)
        print(f"[Warning] formatText failed, added plain text. Error: {error}")
